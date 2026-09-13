from pathlib import Path

import pytest

from app.parsers.document_parser import extract_text

from app.resume.section_detector import detect_section, split_into_sections

from app.resume.contact_extractor import extract_contact_information

from app.resume.content_extractor import (extract_experience,
    extract_education, extract_projects, extract_skills)

from docx import Document

from app.resume.resume_parser import parse_resume

from app.jd.section_detector import (
    detect_section as detect_jd_section,
    split_into_sections as split_jd_sections,
)

FIXTURES_DIR = Path(__file__).parent / "fixtures"


def test_unsupported_file():
    test_file = FIXTURES_DIR / "resume.txt"

    with pytest.raises(ValueError):
        extract_text(test_file)


def test_unsupported_extension():
    test_file = FIXTURES_DIR / "resume.xyz"

    with pytest.raises(ValueError):
        extract_text(test_file)

def test_detect_section():
    assert detect_section("WORK Experience") == "experience"
    assert detect_section("Technical Skills") == "skills"
    assert detect_section("Education") == "education"
    assert detect_section("Projects") == "projects"
    assert detect_section("Random Heading") is None

def test_split_into_sections():
    resume_text = """
    John Doe
    john@example.com

    WORK Experience
    Software Engineer at XYZ Corp
    built and maintained web applications.

    Education
    Bachelor of Science in Computer Science

    Skills
    Python, JavaScript, SQL
    """

    sections = split_into_sections(resume_text)
    assert "John Doe" in sections["contact"]
    assert "Software Engineer at XYZ Corp" in sections["experience"]
    assert "Bachelor of Science in Computer Science" in sections["education"]
    assert "Python, JavaScript, SQL" in sections["skills"]
    assert "built and maintained web applications." in sections["experience"]

def test_extract_contact_information():
    lines = [
        "John Doe",
        "john@example.com",
        "+1 (555) 123-4567",
        "https://www.linkedin.com/in/johndoe",
        "https://github.com/johndoe"
    ]

    contact = extract_contact_information(lines)
    assert contact["name"] == "John Doe"
    assert contact["email"] == "john@example.com"
    assert contact["phone"] == "+1 (555) 123-4567"
    assert contact["linkedin"] == "https://www.linkedin.com/in/johndoe"
    assert contact["github"] == "https://github.com/johndoe"

def test_extract_experience():
    lines = [
        "AI Engineer | Company XYZ",
        "Built machine learning models.",
    ]

    experience = extract_experience(lines)

    assert len(experience) == 2
    assert experience[0]["description"] == "AI Engineer | Company XYZ"
    assert experience[1]["description"] == "Built machine learning models."

def test_extract_education():
    lines = [
        "B.Sc. Artificial Intelligence",
        "Technical University",
    ]

    education = extract_education(lines)

    assert len(education) == 2


def test_extract_projects():
    lines = [
        "ATS Resume Checker",
        "Built an NLP-based resume analysis system.",
    ]

    projects = extract_projects(lines)

    assert len(projects) == 2


def test_extract_skills():
    lines = [
        "Python",
        "PyTorch",
        "Machine Learning",
    ]

    skills = extract_skills(lines)

    assert skills == [
        "Python",
        "PyTorch",
        "Machine Learning",
    ]

def test_parse_resume():
    test_file = FIXTURES_DIR / "test_resume.docx"

    document = Document()

    document.add_paragraph("John Doe")
    document.add_paragraph("john@example.com")
    document.add_paragraph("+49 123 456789")
    document.add_paragraph("linkedin.com/in/johndoe")

    document.add_paragraph("WORK EXPERIENCE")
    document.add_paragraph("AI Engineer | Company XYZ")
    document.add_paragraph("Built machine learning models.")

    document.add_paragraph("EDUCATION")
    document.add_paragraph("B.Sc. Artificial Intelligence")

    document.add_paragraph("PROJECTS")
    document.add_paragraph("ATS Resume Checker")
    document.add_paragraph("Built an NLP-based resume analysis system.")

    document.add_paragraph("SKILLS")
    document.add_paragraph("Python")
    document.add_paragraph("PyTorch")
    document.add_paragraph("Machine Learning")

    document.add_paragraph("ACHIEVEMENTS")
    document.add_paragraph("Winner - Bayerwald Hackathon")

    document.save(test_file)

    resume = parse_resume(test_file)

    assert resume["contact"]["name"] == "John Doe"
    assert resume["contact"]["email"] == "john@example.com"

    assert len(resume["experience"]) == 2
    assert len(resume["education"]) == 1
    assert len(resume["projects"]) == 2
    assert len(resume["achievements"]) == 1
    assert resume["achievements"][0]["description"] == "Winner - Bayerwald Hackathon"

    assert "Python" in resume["skills"]
    assert "PyTorch" in resume["skills"]

    test_file.unlink()

def test_detect_german_sections():
    assert detect_section("BERUFSERFAHRUNG") == "experience"
    assert detect_section("AUSBILDUNG") == "education"
    assert detect_section("PROJEKTE") == "projects"
    assert detect_section("KENNTNISSE") == "skills"
    assert detect_section("ERFOLGE") == "achievements"

def test_detect_additional_german_sections():
    assert detect_section("Praktische Erfahrung") == "experience"
    assert detect_section("Fähigkeiten") == "skills"
    assert detect_section("Erfolge") == "achievements"


def test_detect_jd_sections():
    assert detect_jd_section("Responsibilities") == "responsibilities"
    assert detect_jd_section("Requirements") == "requirements"
    assert detect_jd_section("Education") == "education"
    assert detect_jd_section("Experience") == "experience"
    assert detect_jd_section("Technical Skills") == "skills"


def test_detect_german_jd_sections():
    assert detect_jd_section("Aufgaben") == "responsibilities"
    assert detect_jd_section("Anforderungen") == "requirements"
    assert detect_jd_section("Ausbildung") == "education"
    assert detect_jd_section("Berufserfahrung") == "experience"
    assert detect_jd_section("Kenntnisse") == "skills"


def test_split_jd_sections():
    jd_text = """
    AI Engineer

    RESPONSIBILITIES
    Develop machine learning models.
    Build REST APIs.

    REQUIREMENTS
    Python
    PyTorch

    EDUCATION
    Bachelor's degree in Computer Science.

    EXPERIENCE
    2+ years of experience.

    SKILLS
    FastAPI
    Docker
    """

    sections = split_jd_sections(jd_text)

    assert "AI Engineer" in sections["general"]

    assert "Develop machine learning models." in sections["responsibilities"]
    assert "Build REST APIs." in sections["responsibilities"]

    assert "Python" in sections["requirements"]
    assert "PyTorch" in sections["requirements"]

    assert "Bachelor's degree in Computer Science." in sections["education"]

    assert "2+ years of experience." in sections["experience"]

    assert "FastAPI" in sections["skills"]
    assert "Docker" in sections["skills"]