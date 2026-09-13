from pathlib import Path

from docx import Document


def extract_text_from_docx(file_path: str | Path) -> str:
    """Extract text from paragraphs and tables in a DOCX file."""

    document = Document(file_path)

    content = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            content.append(text)

    for table in document.tables:
        for row in table.rows:
            row_text = []

            for cell in row.cells:
                text = cell.text.strip()

                if text:
                    row_text.append(text)

            if row_text:
                content.append(" | ".join(row_text))

    return "\n".join(content).strip()